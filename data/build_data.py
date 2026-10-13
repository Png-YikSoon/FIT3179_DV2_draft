"""Build chart-ready CSVs for FIT3179 DV2 (domestic air travel in Australia)."""
import csv, math, datetime as dt
from collections import defaultdict
import openpyxl
import pandas as pd

UP = "./raw/"  # folder holding the original BITRE and TRA downloads
ROUTES_X = UP + "toproutesjuly2014july2026.xlsx"
TRA_X = UP + "tra-dots-june-quarter-2026-compared-to-june-quarter-2025.xlsx"
OUT = "./"
MON = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]

def fy(y, m):  # Australian financial year label: Jul 2025-Jun 2026 -> "FY2026"
    return y + 1 if m >= 7 else y

def fy_label(f):
    return f"{f-1}–{str(f)[2:]}"  # e.g. 2025–26

# ---------------------------------------------------------------- BITRE totals
wb = openpyxl.load_workbook(ROUTES_X, read_only=True, data_only=True)
tot = []
for r in wb["Totals"].iter_rows(values_only=True):
    if isinstance(r[0], dt.datetime) and r[1] == "Industry Total":
        d = r[0]
        tot.append(dict(date=d.strftime("%Y-%m-01"), year=d.year, month_num=d.month,
                        month=MON[d.month-1], passengers=r[4], seats=r[9], flights=r[8],
                        load_factor=round(r[7], 1)))
tot = pd.DataFrame(tot)

# ---------------------------------------------------------------- BITRE routes
rows = []
for r in wb["Top Routes"].iter_rows(values_only=True):
    if isinstance(r[2], int) and isinstance(r[3], int) and r[0] and r[1]:
        pax = 0 if r[4] == "COVID19 No flights" else (r[4] if isinstance(r[4], (int, float)) else None)
        rows.append(dict(a=r[0], b=r[1], year=r[2], month_num=r[3], passengers=pax,
                         flights=r[5], load_factor=r[6], distance_km=r[7], seats=r[10],
                         name=r[20]))
rt = pd.DataFrame(rows)
rt["passengers"] = pd.to_numeric(rt.passengers)
rt["route_code"] = rt.a + "-" + rt.b

# airport code -> city name, from the "City A - City B" label
# BITRE labels don't always follow code order, so each code's city is the name
# common to every route label it appears in (resolved by elimination).
cand = {}
pairs = rt.drop_duplicates("route_code")[["a","b","name"]].values.tolist()
for a, b, n in pairs:
    names = {p.strip() for p in str(n).split(" - ")}
    for c in (a, b):
        cand[c] = cand[c] & names if c in cand else set(names)
city = {}
for _ in range(10):
    for c, sset in cand.items():
        if len(sset) == 1: city[c] = next(iter(sset))
    for a, b, n in pairs:
        names = {p.strip() for p in str(n).split(" - ")}
        for x, y in ((a, b), (b, a)):
            if x in city and y not in city and len(names - {city[x]}) == 1:
                city[y] = next(iter(names - {city[x]}))
    for c in city: cand[c] = {city[c]}
assert len(city) == len(cand), set(cand) - set(city)
assert len(set(city.values())) == len(city), "duplicate city names"
rt["route"] = rt.a.map(city) + " – " + rt.b.map(city)
rt["fy"] = [fy(y, m) for y, m in zip(rt.year, rt.month_num)]

# airport coordinates (OurAirports, public domain)
ap = pd.read_csv("./raw/ourairports_airports.csv", low_memory=False)
ap = ap[(ap.iso_country == "AU") & ap.iata_code.notna()].drop_duplicates("iata_code")
ap = ap.set_index("iata_code")
STATE = {"AU-NSW":"NSW","AU-VIC":"VIC","AU-QLD":"QLD","AU-SA":"SA","AU-WA":"WA",
         "AU-TAS":"TAS","AU-NT":"NT","AU-ACT":"ACT"}
ALIAS = {"QNA": "BNK"}  # BITRE uses QNA for Ballina Byron Gateway (IATA BNK)
codes = sorted(city)
missing = [c for c in codes if ALIAS.get(c, c) not in ap.index]
assert not missing, missing
airports = pd.DataFrame(dict(
    code=codes, city=[city[c] for c in codes],
    state=[STATE.get(ap.loc[ALIAS.get(c, c), "iso_region"], "") for c in codes],
    lat=[round(ap.loc[ALIAS.get(c, c), "latitude_deg"], 4) for c in codes],
    lon=[round(ap.loc[ALIAS.get(c, c), "longitude_deg"], 4) for c in codes]))
airports.to_csv(OUT + "airports.csv", index=False)
A = airports.set_index("code")

# ------------------------------------------------------- 01 long-term trend
c1 = tot[["date","year","month","passengers","seats","flights","load_factor"]].copy()
c1["passengers_12m_avg"] = c1.passengers.rolling(12).mean().round(0)
c1.to_csv(OUT + "01_trend_monthly_passengers.csv", index=False)

# ------------------------------------------------------- 02 radial calendar
c2 = tot[["year","month_num","month","passengers"]].copy()
c2["passengers_m"] = (c2.passengers / 1e6).round(3)
# index vs that calendar year's own monthly average (shows seasonality, not growth)
c2["index_vs_year_avg"] = (c2.passengers / c2.groupby("year").passengers.transform("mean") * 100).round(1)
c2["complete_year"] = c2.groupby("year").month_num.transform("count").eq(12)
c2.loc[~c2.complete_year, "index_vs_year_avg"] = None  # 2014 and 2026 are partial years
c2.to_csv(OUT + "02_calendar_monthly.csv", index=False)

# ------------------------------------------------------- 03 flow map
LATEST, BASE = 2026, 2019  # FY2025–26 vs FY2018–19
g = rt.groupby(["route_code","fy"]).agg(passengers=("passengers","sum"),
                                         months=("passengers","count")).reset_index()
fyp = g.pivot(index="route_code", columns="fy", values="passengers")
fym = g.pivot(index="route_code", columns="fy", values="months")
meta = rt.drop_duplicates("route_code").set_index("route_code")[["a","b","route","distance_km"]]
flow = meta.join(fyp[[LATEST]].rename(columns={LATEST:"passengers_fy2026"}))
flow = flow[fym[LATEST].eq(12)].dropna(subset=["passengers_fy2026"])
for s, c in (("origin","a"),("dest","b")):
    flow[f"{s}_code"] = flow[c]; flow[f"{s}_city"] = flow[c].map(A.city)
    flow[f"{s}_lat"] = flow[c].map(A.lat); flow[f"{s}_lon"] = flow[c].map(A.lon)
flow = flow.reset_index().sort_values("passengers_fy2026", ascending=False)
flow["rank"] = range(1, len(flow)+1)
flow[["rank","route_code","route","origin_code","origin_city","origin_lat","origin_lon",
      "dest_code","dest_city","dest_lat","dest_lon","distance_km","passengers_fy2026"]
     ].to_csv(OUT + "03_flowmap_routes_fy2026.csv", index=False)

# ------------------------------------------------------- 04 airport symbols
def airport_totals(f):
    full = fym[f].eq(12)
    sub = g[(g.fy == f) & g.route_code.isin(full[full].index)]
    sub = sub.merge(meta[["a","b"]], left_on="route_code", right_index=True)
    return pd.concat([sub[["a","passengers"]].rename(columns={"a":"code"}),
                      sub[["b","passengers"]].rename(columns={"b":"code"})]).groupby("code").passengers.sum()
a26, a19 = airport_totals(LATEST), airport_totals(BASE)
c4 = airports.set_index("code").join(a19.rename("passengers_fy2019")).join(a26.rename("passengers_fy2026"))
c4 = c4.dropna(subset=["passengers_fy2026"]).reset_index()
c4["recovery_pct_of_2019"] = (c4.passengers_fy2026 / c4.passengers_fy2019 * 100).round(1)
c4["change_pct"] = (c4.recovery_pct_of_2019 - 100).round(1)
c4.sort_values("passengers_fy2026", ascending=False).to_csv(OUT + "04_airports_traffic_recovery.csv", index=False)

# ------------------------------------------------------- 05 bump chart
full_fy = sorted(f for f in g.fy.unique() if g[g.fy == f].months.max() == 12 and f <= LATEST)
bump = []
for f in full_fy:
    s = g[(g.fy == f) & (g.months == 12)].sort_values("passengers", ascending=False).head(10)
    for i, r in enumerate(s.itertuples(), 1):
        bump.append(dict(fy=f, fy_label=fy_label(f), route_code=r.route_code,
                         route=meta.loc[r.route_code, "route"], passengers=int(r.passengers), rank=i,
                         covid_data_gap=f in (2020, 2021, 2022)))
pd.DataFrame(bump).to_csv(OUT + "05_bump_top10_routes_by_fy.csv", index=False)

# ------------------------------------------------------- 06 load-factor heatmap
top15 = flow.head(15).route_code.tolist()
h = rt[(rt.fy == LATEST) & rt.route_code.isin(top15)].copy()
h["load_factor"] = h.load_factor.round(1)
h["month"] = h.month_num.map(lambda m: MON[m-1])
h["fy_month_order"] = (h.month_num - 7) % 12 + 1  # Jul=1 ... Jun=12
h["route_rank"] = h.route_code.map({c: i for i, c in enumerate(top15, 1)})
h[["route_rank","route_code","route","year","month_num","month","fy_month_order",
   "passengers","seats","load_factor"]].sort_values(["route_rank","fy_month_order"]
   ).to_csv(OUT + "06_heatmap_loadfactor_top15_fy2026.csv", index=False)

# ------------------------------------------------------- 10 slopegraph
cy = rt.groupby(["route_code","year"]).agg(p=("passengers","sum"), n=("passengers","count")).reset_index()
cyp = cy[cy.n == 12].pivot(index="route_code", columns="year", values="p")
sl = cyp[[2019, 2025]].dropna()
sl = sl.loc[sl[2019].sort_values(ascending=False).head(15).index]
sl = meta[["route"]].join(sl, how="inner").reset_index()
sl.columns = ["route_code","route","passengers_2019","passengers_2025"]
sl["change_pct"] = ((sl.passengers_2025 / sl.passengers_2019 - 1) * 100).round(1)
sl["recovered"] = sl.passengers_2025 >= sl.passengers_2019
long = sl.melt(id_vars=["route_code","route","change_pct","recovered"],
               value_vars=["passengers_2019","passengers_2025"], var_name="year", value_name="passengers")
long["year"] = long.year.str[-4:].astype(int)
long.sort_values(["route_code","year"]).to_csv(OUT + "10_slopegraph_routes_2019_vs_2025.csv", index=False)

# ------------------------------------------------------- 11 seasonal small multiples
# mix of business corridors, leisure routes and fly-in fly-out (FIFO) mining routes
SEL = {"MEL-SYD":"Business corridor","CBR-SYD":"Business corridor","BNE-SYD":"Business corridor",
       "MEL-OOL":"Leisure","CNS-SYD":"Leisure","HBA-MEL":"Leisure",
       "PER-PHE":"Mining (FIFO)","KGI-PER":"Mining (FIFO)"}
s = rt[rt.route_code.isin(list(SEL)) & rt.year.between(2023, 2025)]
s = s.groupby(["route_code","route","month_num"]).passengers.mean().reset_index()
s["month"] = s.month_num.map(lambda m: MON[m-1])
s["route_type"] = s.route_code.map(SEL)
s["avg_passengers"] = s.passengers.round(0)
s["index_vs_route_avg"] = (s.passengers / s.groupby("route_code").passengers.transform("mean") * 100).round(1)
s[["route_type","route_code","route","month_num","month","avg_passengers","index_vs_route_avg"]
  ].to_csv(OUT + "11_seasonal_selected_routes_2023_2025.csv", index=False)

# ---------------------------------------------------------------- TRA helpers
tw = openpyxl.load_workbook(TRA_X, read_only=True, data_only=True)
def block(sheet, heading):
    """Return {label: [values...]} for the block whose first cell starts with heading."""
    out, on = {}, False
    for r in tw[sheet].iter_rows(values_only=True):
        r = [c for c in r if c is not None]
        if not r: continue
        if isinstance(r[0], str) and r[0].startswith(heading): on = True; continue
        if on:
            if len(r) == 1 or not isinstance(r[1], (int, float)): break
            out[r[0]] = r[1:]
    return out

# ------------------------------------------------------- 07 choropleth
ovn = block("Table 1c", "Overnight trips"); spend = block("Table 1c", "Regional expenditure")
alltrips = block("Table 9c", "Total domestic trips")
c7 = []
for st, ab in [("New South Wales","NSW"),("Victoria","VIC"),("Queensland","QLD"),("South Australia","SA"),
               ("Western Australia","WA"),("Tasmania","TAS"),("Northern Territory","NT"),
               ("Australian Capital Territory","ACT")]:
    c7.append(dict(state=st, state_abbr=ab,
        overnight_trips_000_ye_jun2025=round(ovn[st][0],1), overnight_trips_000_ye_jun2026=round(ovn[st][1],1),
        overnight_trips_change_pct=round(ovn[st][3],1),
        overnight_spend_m_ye_jun2025=round(spend[st][0],1), overnight_spend_m_ye_jun2026=round(spend[st][1],1),
        overnight_spend_change_pct=round(spend[st][3],1),
        all_trips_000_ye_jun2025=round(alltrips[st][0],1), all_trips_000_ye_jun2026=round(alltrips[st][1],1),
        all_trips_change_pct=round(alltrips[st][3],1),
        small_sample_caution=ab in ("TAS","NT","ACT")))
pd.DataFrame(c7).to_csv(OUT + "07_choropleth_states_ye_jun2026.csv", index=False)

# ------------------------------------------------------- 08 waffle
P = ["Holiday","Visiting friends and relatives","Business","Other"]
LBL = {"Holiday":"Holiday","Visiting friends and relatives":"Visiting friends & relatives",
       "Business":"Business","Other":"Other"}
def waffle(sheet, heading, trip_type):
    b = block(sheet, heading); v = {p: b[p][-1] for p in P}; tot_ = sum(v.values())
    share = {p: v[p] / tot_ * 100 for p in P}
    # largest-remainder rounding so squares add to exactly 100
    sq = {p: math.floor(share[p]) for p in P}
    for p in sorted(P, key=lambda p: share[p] - sq[p], reverse=True)[:100 - sum(sq.values())]:
        sq[p] += 1
    rows, cells = [], []
    for p in P:
        rows.append(dict(trip_type=trip_type, purpose=LBL[p], trips_000=round(v[p],1),
                         share_pct=round(share[p],1), squares=sq[p]))
        cells += [LBL[p]] * sq[p]
    grid = [dict(trip_type=trip_type, cell=i+1, col=i % 10 + 1, row=i // 10 + 1, purpose=c)
            for i, c in enumerate(cells)]
    return rows, grid
r1, g1 = waffle("Table 11c", "Total domestic trips", "All trips (day + overnight)")
r2, g2 = waffle("Table 3c", "Overnight trips", "Overnight trips")
pd.DataFrame(r1 + r2).to_csv(OUT + "08_waffle_purpose_summary_ye_jun2026.csv", index=False)
pd.DataFrame(g1 + g2).to_csv(OUT + "08_waffle_purpose_grid_ye_jun2026.csv", index=False)

# ------------------------------------------------------- 09 purpose over quarters
Q = ["Jun Q 2025","Sep Q 2025","Dec Q 2025","Mar Q 2026","Jun Q 2026"]
c9 = []
for sheet, head, tt in [("Table 11b","Total domestic trips","All trips"),
                        ("Table 3b","Overnight trips","Overnight trips")]:
    b = block(sheet, head)
    for p in P:
        for i, q in enumerate(Q):
            c9.append(dict(trip_type=tt, purpose=LBL[p], quarter=q, quarter_order=i+1,
                           trips_000=round(b[p][i],1),
                           index_vs_jun_q_2025=round(b[p][i]/b[p][0]*100,1)))
pd.DataFrame(c9).to_csv(OUT + "09_purpose_by_quarter.csv", index=False)
air = []
for sheet, tt in [("Table 4a","Overnight trips"),("Table 8a","Day trips"),("Table 12a","All trips")]:
    for r in tw[sheet].iter_rows(values_only=True):
        r = [c for c in r if c is not None]
        if r and r[0] == "Domestic airfares":
            for i, q in enumerate(Q):
                air.append(dict(trip_type=tt, quarter=q, quarter_order=i+1, airfare_spend_m=round(r[1+i],1)))
pd.DataFrame(air).to_csv(OUT + "09b_airfare_spend_by_quarter.csv", index=False)

# ------------------------------------------------------- 12 dumbbell (TRA PDF, Figure 2)
d = [("1 night",17,781,23,2000),("2–3 nights",50,2300,48,4000),
     ("4–7 nights",27,1200,21,1800),("8+ nights",6,282,8,637)]
c12 = []
for i, (k, pc, tc, pn, tn) in enumerate(d, 1):
    c12 += [dict(trip_length=k, order=i, party="With children", share_pct=pc, trips_000=tc),
            dict(trip_length=k, order=i, party="No children", share_pct=pn, trips_000=tn)]
pd.DataFrame(c12).to_csv(OUT + "12_dumbbell_holiday_length_children_jun_q_2026.csv", index=False)
print("done")
